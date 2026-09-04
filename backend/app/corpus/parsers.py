from __future__ import annotations

import asyncio
from concurrent.futures import ThreadPoolExecutor
import logging
import re
from collections import Counter
from pathlib import Path

from app.corpus.models import ParsedBlock, ParsedDocument, SourceEntry
from app.corpus.file_detection import FileTypeDetectionError, PdfPreflight, preflight_pdf, render_pdf_pages
from app.corpus.visual import enrich_document_visuals
from app.core.config import get_settings
from app.services.mineru_service import MinerUParser, MinerUStructuredResult


logger = logging.getLogger(__name__)


class CorpusParseError(ValueError):
    pass


MOJIBAKE_MARKERS = ("锟斤拷", "ï¿½", "Ã", "Â", "�")


def clean_text(text: str) -> tuple[str, list[str]]:
    raw = str(text or "").replace("\r\n", "\n").replace("\r", "\n")
    warnings: list[str] = []
    cleaned = re.sub(r"[\u0000-\u0008\u000b\u000c\u000e-\u001f\u007f]", " ", raw)
    if cleaned != raw:
        warnings.append("control_characters_removed")
    cleaned = re.sub(r"[ \t]+", " ", cleaned)
    cleaned = re.sub(r" *\n *", "\n", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
    if any(marker in cleaned for marker in MOJIBAKE_MARKERS):
        warnings.append("possible_mojibake")
    return cleaned, warnings


def _remove_repeated_page_edges(pages: list[str]) -> tuple[list[str], list[str]]:
    if len(pages) < 3:
        return pages, []
    edges: list[str] = []
    page_lines: list[list[str]] = []
    for page in pages:
        lines = [line.strip() for line in page.splitlines() if line.strip()]
        page_lines.append(lines)
        if lines:
            edges.extend([lines[0], lines[-1]])
    threshold = max(3, int(len(pages) * 0.6))
    repeated = {line for line, count in Counter(edges).items() if count >= threshold and len(line) < 120}
    if not repeated:
        return pages, []
    output = ["\n".join(line for line in lines if line not in repeated) for lines in page_lines]
    return output, ["repeated_headers_or_footers_removed"]


def _blocks_from_text(text: str, *, page: int | None = None) -> list[ParsedBlock]:
    blocks: list[ParsedBlock] = []
    section: list[str] = []
    cursor = 0
    for raw in re.split(r"\n\s*\n", text):
        value = raw.strip()
        if not value:
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", value)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).strip()
            section = section[: level - 1] + [title]
            kind = "heading"
            block_text = title
            heading_level = level
        else:
            if value.startswith("```") and value.endswith("```"):
                kind = "code"
            elif value.startswith("$$") and value.endswith("$$"):
                kind = "formula"
            elif re.search(r"!\[[^]]*]\([^)]+\)", value):
                kind = "image"
            elif len(value.splitlines()) >= 2 and all("|" in line for line in value.splitlines()):
                kind = "table"
            elif all(re.match(r"^(?:[-*+] |\d+[.)] )", line.strip()) for line in value.splitlines() if line.strip()):
                kind = "list"
            else:
                kind = "paragraph"
            block_text = value
            heading_level = None
        start = text.find(raw, cursor)
        cursor = max(cursor, start + len(raw))
        blocks.append(
            ParsedBlock(
                text=block_text,
                kind=kind,
                page_start=page,
                page_end=page,
                section_path=list(section),
                start_index=max(0, start),
                heading_level=heading_level,
            )
        )
    return blocks


def _parse_markdown(path: Path, source: SourceEntry) -> ParsedDocument:
    raw = path.read_text(encoding="utf-8", errors="replace")
    text, warnings = clean_text(raw)
    if not text:
        raise CorpusParseError("empty_document")
    return ParsedDocument(source=source, text=text, blocks=_blocks_from_text(text), warnings=warnings)


def _parse_docx(path: Path, source: SourceEntry) -> ParsedDocument:
    try:
        from docx import Document
        document = Document(str(path))
        blocks: list[ParsedBlock] = []
        section: list[str] = []
        cursor = 0
        for paragraph in document.paragraphs:
            value, paragraph_warnings = clean_text(paragraph.text)
            if not value:
                continue
            style = str(getattr(paragraph.style, "name", "") or "")
            heading_match = re.search(r"(?:heading|标题)\s*([1-6])", style, flags=re.IGNORECASE)
            if heading_match:
                level = int(heading_match.group(1))
                section = section[: level - 1] + [value]
                kind = "heading"
            elif "list" in style.lower() or style.startswith("列表"):
                level = None
                kind = "list"
            else:
                level = None
                kind = "paragraph"
            blocks.append(ParsedBlock(value, kind, section_path=list(section), start_index=cursor, heading_level=level))
            cursor += len(value) + 2
        for table_index, table in enumerate(document.tables):
            rows = []
            for row in table.rows:
                cells = [clean_text(cell.text)[0] for cell in row.cells]
                rows.append("<tr>" + "".join(f"<td>{cell}</td>" for cell in cells) + "</tr>")
            if rows:
                value = "<table>" + "".join(rows) + "</table>"
                blocks.append(ParsedBlock(value, "table", section_path=list(section), start_index=cursor, metadata={"source_table_index": table_index}))
                cursor += len(value) + 2
        # Headers and footers are retained as explicit source blocks so they
        # can be filtered later without losing provenance.
        for section_index, document_section in enumerate(document.sections):
            for label, container in (("header", document_section.header), ("footer", document_section.footer)):
                value = clean_text("\n".join(item.text for item in container.paragraphs))[0]
                if value:
                    blocks.append(ParsedBlock(value, label, section_path=list(section), start_index=cursor, metadata={"section_index": section_index}))
                    cursor += len(value) + 2
        text = "\n\n".join(block.text for block in blocks if block.kind not in {"heading"})
        if not text:
            raise CorpusParseError("empty_document")
        return ParsedDocument(source=source, text=text, blocks=blocks, parser_schema="docx")
    except Exception as exc:
        # Keep the existing docx2txt fallback for minimal/legacy OOXML files.
        try:
            import docx2txt
            raw = docx2txt.process(str(path)) or ""
        except Exception as fallback_exc:
            raise CorpusParseError(f"unreadable_docx:{type(fallback_exc).__name__}") from exc
        text, warnings = clean_text(raw)
        if not text:
            raise CorpusParseError("empty_document")
        return ParsedDocument(source=source, text=text, blocks=_blocks_from_text(text), warnings=warnings)


def _parse_legacy_doc(path: Path, source: SourceEntry) -> ParsedDocument:
    """Read the RTF payload commonly exported with a .doc extension.

    Binary OLE .doc files still go to MinerU; this local fallback only handles
    the RTF form so an upstream outage does not make the supplied Word corpus
    unusable.
    """
    raw = path.read_bytes()
    if not raw.lstrip().startswith(b"{\\rtf"):
        raise CorpusParseError("legacy_doc_requires_mineru")
    text = raw.decode("latin-1", errors="replace")

    # RTF is a group-based format. A simple control-word regex leaves shape
    # and font-table payloads in the output, so walk the stream and skip
    # ignorable destinations (fonttbl, stylesheet, shp, pict, ...).
    destination_names = {"fonttbl", "stylesheet", "colortbl", "info", "pict", "shp", "shpinst", "listtable", "listoverridetable"}
    output: list[str] = []
    stack: list[bool] = [False]
    token_re = re.compile(r"[{}]|\\(?:'(?P<byte>[0-9a-fA-F]{2})|(?P<word>[a-zA-Z]+)(?P<num>-?\d+)? ?|(?P<symbol>[{}\\]))")
    cursor = 0
    for match in token_re.finditer(text):
        if not stack[-1] and match.start() > cursor:
            output.append(text[cursor : match.start()])
        token = match.group(0)
        if token == "{":
            # A destination control word immediately following an opening
            # group marks the whole group as non-body metadata.
            tail = text[match.end() : match.end() + 32]
            destination = re.match(r"\\([a-zA-Z]+|\*)", tail)
            stack.append(stack[-1] or bool(destination and destination.group(1).lower() in destination_names | {"*"}))
        elif token == "}":
            if len(stack) > 1:
                stack.pop()
        elif not stack[-1]:
            if match.group("byte"):
                output.append(bytes.fromhex(match.group("byte")).decode("cp936", errors="replace"))
            elif match.group("symbol"):
                output.append(match.group("symbol"))
            else:
                word = (match.group("word") or "").lower()
                number = match.group("num")
                if word == "u" and number is not None:
                    codepoint = int(number)
                    output.append(chr(codepoint if codepoint >= 0 else codepoint + 65536))
                elif word in {"par", "line", "row", "cell", "sect"}:
                    output.append("\n")
                elif word == "tab":
                    output.append("\t")
        cursor = match.end()
    if not stack[-1] and cursor < len(text):
        output.append(text[cursor:])
    text = "".join(output)
    text = re.sub(r"[;]+", " ", text)
    text, warnings = clean_text(text)
    if not text:
        raise CorpusParseError("empty_document")
    return ParsedDocument(source=source, text=text, blocks=_blocks_from_text(text), warnings=warnings, parser_schema="rtf")


def _parse_image(path: Path, source: SourceEntry) -> ParsedDocument:
    try:
        from PIL import Image
        with Image.open(path) as image:
            image.verify()
    except Exception as exc:
        raise CorpusParseError(f"unreadable_image:{type(exc).__name__}") from exc
    block = ParsedBlock(
        text=f"[image: {path.name}]",
        kind="image",
        media_path=str(path),
        metadata={"visual_status": "pending", "source_type": "image"},
    )
    return ParsedDocument(source=source, text=block.text, blocks=[block], parser_schema="image")


def _parse_rendered_pdf_pages(source: SourceEntry, page_paths: list[str], preflight: PdfPreflight) -> ParsedDocument:
    """Create traceable image blocks when remote OCR is temporarily unavailable."""
    blocks = [
        ParsedBlock(
            text=f"[scanned page {index}]",
            kind="image",
            page_start=index,
            page_end=index,
            media_path=path,
            metadata={"source_type": "scanned_pdf_page", "ocr_status": "degraded", "page_number": index},
        )
        for index, path in enumerate(page_paths, start=1)
    ]
    return ParsedDocument(
        source=source,
        text="\n\n".join(block.text for block in blocks),
        blocks=blocks,
        page_count=preflight.page_count,
        warnings=[*preflight.warnings, "ocr_unavailable_degraded"],
        parser_name="pdf-page-render-fallback",
        parser_schema="scanned-pages",
        metadata={
            "pdf_preflight": {
                "content_profile": preflight.content_profile,
                "page_count": preflight.page_count,
                "empty_text_pages": preflight.empty_text_pages,
                "text_density": preflight.text_density,
                "non_empty_page_ratio": preflight.non_empty_page_ratio,
                "ocr_required": preflight.ocr_required,
                "warnings": list(preflight.warnings),
                "rendered_page_paths": list(page_paths),
                "pages": [page.__dict__ for page in preflight.pages],
            }
        },
    )


def _parse_pdf_with_rendered_fallback(
    source: SourceEntry,
    path: Path,
    page_paths: list[str],
    preflight: PdfPreflight,
) -> ParsedDocument:
    """Keep text pages and attach degraded image blocks for scanned pages."""
    if preflight.content_profile == "scanned_pdf":
        return _parse_rendered_pdf_pages(source, page_paths, preflight)
    try:
        parsed = _parse_pdf(path, source)
    except CorpusParseError as exc:
        if str(exc) != "scanned_or_empty_pdf":
            raise
        return _parse_rendered_pdf_pages(source, page_paths, preflight)

    existing_pages = {block.page_start for block in parsed.blocks if block.page_start is not None}
    image_blocks: list[ParsedBlock] = []
    for index, page_path in enumerate(page_paths, start=1):
        profile = preflight.pages[index - 1].profile if index - 1 < len(preflight.pages) else "scanned_page"
        if profile == "text_page" or index in existing_pages:
            continue
        image_blocks.append(
            ParsedBlock(
                text=f"[scanned page {index}]",
                kind="image",
                page_start=index,
                page_end=index,
                media_path=page_path,
                metadata={"source_type": "scanned_pdf_page", "ocr_status": "degraded", "page_number": index},
            )
        )
    if not image_blocks:
        return _attach_pdf_preflight(parsed, preflight)
    blocks = [*parsed.blocks, *image_blocks]
    blocks.sort(key=lambda block: (block.page_start or 0, block.start_index))
    return replace(
        _attach_pdf_preflight(parsed, preflight),
        blocks=blocks,
        text=parsed.text + "\n\n" + "\n\n".join(block.text for block in image_blocks),
        parser_name="pdf-page-render-fallback",
        parser_schema="mixed-pages",
        warnings=list(dict.fromkeys([*parsed.warnings, "ocr_unavailable_degraded"])),
    )


def _parse_pdf(path: Path, source: SourceEntry) -> ParsedDocument:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise CorpusParseError("encrypted_pdf")
        raw_pages = [(page.extract_text() or "") for page in reader.pages]
    except CorpusParseError:
        raise
    except Exception as exc:
        raise CorpusParseError(f"unreadable_pdf:{type(exc).__name__}") from exc
    pages, edge_warnings = _remove_repeated_page_edges(raw_pages)
    cleaned_pages: list[str] = []
    warnings = list(edge_warnings)
    for page in pages:
        cleaned, page_warnings = clean_text(page)
        cleaned_pages.append(cleaned)
        warnings.extend(page_warnings)
    if not any(cleaned_pages):
        raise CorpusParseError("scanned_or_empty_pdf")
    blocks: list[ParsedBlock] = []
    offset = 0
    for page_index, page_text in enumerate(cleaned_pages, start=1):
        for block in _blocks_from_text(page_text, page=page_index):
            blocks.append(
                ParsedBlock(
                    text=block.text,
                    kind=block.kind,
                    page_start=page_index,
                    page_end=page_index,
                    section_path=block.section_path,
                    start_index=offset + block.start_index,
                    heading_level=block.heading_level,
                    bbox=block.bbox,
                    media_path=block.media_path,
                    mineru_node_path=block.mineru_node_path,
                    metadata=block.metadata,
                )
            )
        offset += len(page_text) + 2
    text = "\n\n".join(page for page in cleaned_pages if page)
    return ParsedDocument(
        source=source,
        text=text,
        blocks=blocks,
        page_count=len(raw_pages),
        warnings=list(dict.fromkeys(warnings)),
    )


def parsed_document_from_mineru(
    result: MinerUStructuredResult,
    source: SourceEntry,
    *,
    preflight: PdfPreflight | None = None,
) -> ParsedDocument:
    metadata = {"markdown_available": bool(result.markdown)}
    if preflight is not None:
        metadata["pdf_preflight"] = {
            "content_profile": preflight.content_profile,
            "page_count": preflight.page_count,
            "empty_text_pages": preflight.empty_text_pages,
            "text_density": preflight.text_density,
            "non_empty_page_ratio": preflight.non_empty_page_ratio,
            "ocr_required": preflight.ocr_required,
            "pages": [page.__dict__ for page in preflight.pages],
        }
    return ParsedDocument(
        source=source,
        text=result.text,
        blocks=[
            ParsedBlock(
                text=block.text,
                kind=block.kind,
                page_start=block.page_start,
                page_end=block.page_end,
                section_path=list(block.section_path),
                start_index=block.start_index,
                heading_level=block.heading_level,
                bbox=list(block.bbox) if block.bbox else None,
                media_path=block.media_path,
                mineru_node_path=block.node_path,
                metadata=dict(block.metadata),
            )
            for block in result.blocks
        ],
        page_count=result.page_count,
        warnings=list(result.warnings),
        parser_name="mineru",
        parser_schema=result.schema_version,
        metadata=metadata,
    )


def _attach_pdf_preflight(parsed: ParsedDocument, preflight: PdfPreflight | None) -> ParsedDocument:
    if preflight is None:
        return parsed
    metadata = dict(parsed.metadata)
    metadata["pdf_preflight"] = {
        "content_profile": preflight.content_profile,
        "page_count": preflight.page_count,
        "empty_text_pages": preflight.empty_text_pages,
        "text_density": preflight.text_density,
        "non_empty_page_ratio": preflight.non_empty_page_ratio,
        "ocr_required": preflight.ocr_required,
        "warnings": list(preflight.warnings),
        "pages": [page.__dict__ for page in preflight.pages],
    }
    warnings = list(dict.fromkeys([*parsed.warnings, *preflight.warnings]))
    return ParsedDocument(
        source=parsed.source,
        text=parsed.text,
        blocks=parsed.blocks,
        page_count=parsed.page_count,
        warnings=warnings,
        rejected_blocks=parsed.rejected_blocks,
        parser_name=parsed.parser_name,
        parser_schema=parsed.parser_schema,
        metadata=metadata,
    )


def _parse_source_locally(source: SourceEntry) -> ParsedDocument:
    path = source.resolved_path
    if path is None:
        raise CorpusParseError("remote_source_not_cached")
    if not path.exists() or not path.is_file():
        raise CorpusParseError("source_file_not_found")
    if source.content_type == "pdf":
        return _parse_pdf(path, source)
    if source.content_type == "docx":
        return _parse_docx(path, source)
    if source.content_type == "doc":
        return _parse_legacy_doc(path, source)
    if source.content_type == "markdown":
        return _parse_markdown(path, source)
    if source.content_type == "image":
        return _parse_image(path, source)
    raise CorpusParseError("remote_source_not_cached")


async def parse_source_async(source: SourceEntry) -> ParsedDocument:
    path = source.resolved_path
    if path is None:
        raise CorpusParseError("remote_source_not_cached")
    if not path.exists() or not path.is_file():
        raise CorpusParseError("source_file_not_found")

    settings = get_settings()
    mineru_types = {"pdf", "doc", "docx", "ppt", "pptx", "xls", "xlsx", "html", "image"}
    pdf_preflight: PdfPreflight | None = None
    rendered_page_paths: list[str] = []
    if source.content_type == "pdf":
        # Legacy callers may pass an already-ingested fixture that is not a
        # complete PDF (for example when MinerU is mocked). Only run the
        # read-only preflight for data carrying the PDF signature; upload
        # validation remains the authoritative place for rejecting fakes.
        if path.read_bytes()[:5] == b"%PDF-":
            try:
                pdf_preflight = preflight_pdf(path)
                if pdf_preflight.content_profile == "invalid_pdf":
                    raise CorpusParseError("encrypted_pdf")
                if pdf_preflight.ocr_required:
                    artifact_root = Path(settings.mineru_artifact_dir) / source.document_id / "pages"
                    try:
                        rendered_page_paths = [str(item) for item in render_pdf_pages(path, artifact_root)]
                    except FileTypeDetectionError as exc:
                        logger.warning(
                            "pdf_page_render_failed document_id=%s failure_type=%s",
                            source.document_id,
                            type(exc).__name__,
                        )
                        pdf_preflight = PdfPreflight(
                            pdf_preflight.content_profile,
                            pdf_preflight.page_count,
                            pdf_preflight.pages,
                            pdf_preflight.empty_text_pages,
                            pdf_preflight.text_density,
                            warnings=[*pdf_preflight.warnings, str(exc)],
                        )
            except FileTypeDetectionError as exc:
                # Keep the established parser error contract for malformed
                # PDFs; callers can surface ``unreadable_pdf`` consistently.
                raise CorpusParseError("unreadable_pdf") from exc
    if settings.mineru_enabled and settings.mineru_api_token and source.content_type in mineru_types:
        try:
            result = await MinerUParser().parse_structured(
                str(path),
                document_id=source.document_id,
                original_name=path.name,
                is_ocr=pdf_preflight.ocr_required if pdf_preflight is not None else None,
            )
        except Exception as exc:
            logger.warning(
                "public_mineru_parse_failed document_id=%s failure_type=%s",
                source.document_id,
                type(exc).__name__,
            )
            if source.content_type not in {"pdf", "doc", "docx", "image"}:
                raise CorpusParseError(f"mineru_failed:{type(exc).__name__}") from exc
            if (
                source.content_type == "pdf"
                and pdf_preflight is not None
                and pdf_preflight.ocr_required
                and rendered_page_paths
            ):
                parsed = _parse_rendered_pdf_pages(source, rendered_page_paths, pdf_preflight)
            else:
                parsed = _attach_pdf_preflight(_parse_source_locally(source), pdf_preflight)
            if rendered_page_paths and "pdf_preflight" in parsed.metadata:
                parsed.metadata["pdf_preflight"]["rendered_page_paths"] = rendered_page_paths
            parsed.warnings.append(f"mineru_unavailable_local_fallback:{type(exc).__name__}")
            return await enrich_document_visuals(
                parsed,
                path,
                document_id=source.document_id,
                media_root=settings.mineru_artifact_dir,
            )
        parsed = parsed_document_from_mineru(result, source, preflight=pdf_preflight)
        if rendered_page_paths:
            parsed.metadata["pdf_preflight"]["rendered_page_paths"] = rendered_page_paths
        # Visual enrichment is deliberately outside the MinerU fallback
        # boundary: a Qwen-VL timeout or malformed response must not discard a
        # valid MinerU structured document. enrich_document_visuals itself
        # marks individual assets as degraded and continues.
        try:
            return await enrich_document_visuals(
                parsed,
                path,
                document_id=source.document_id,
                media_root=settings.mineru_artifact_dir,
            )
        except Exception as exc:
            logger.warning("visual_enrichment_skipped document_id=%s failure_type=%s", source.document_id, type(exc).__name__)
            parsed.warnings.append("visual_analysis_degraded")
            return parsed
    if source.content_type in mineru_types and source.content_type not in {"pdf", "doc", "docx", "image"}:
        reason = "disabled" if not settings.mineru_enabled else "token_missing"
        raise CorpusParseError(f"mineru_unavailable:{reason}")
    if (
        source.content_type == "pdf"
        and pdf_preflight is not None
        and pdf_preflight.ocr_required
        and rendered_page_paths
    ):
        parsed = _parse_pdf_with_rendered_fallback(source, path, rendered_page_paths, pdf_preflight)
    else:
        parsed = _attach_pdf_preflight(_parse_source_locally(source), pdf_preflight)
    if rendered_page_paths and "pdf_preflight" in parsed.metadata:
        parsed.metadata["pdf_preflight"]["rendered_page_paths"] = rendered_page_paths
    if source.content_type in {"pdf", "doc", "docx"} and (not settings.mineru_enabled or not settings.mineru_api_token):
        reason = "disabled" if not settings.mineru_enabled else "token_missing"
        parsed.warnings.append(f"mineru_unavailable_local_fallback:{reason}")
    return await enrich_document_visuals(
        parsed,
        path,
        document_id=source.document_id,
        media_root=settings.mineru_artifact_dir,
    )


def parse_source(source: SourceEntry) -> ParsedDocument:
    """Synchronous corpus entry point, safe even when called under an event loop."""
    try:
        asyncio.get_running_loop()
    except RuntimeError:
        return asyncio.run(parse_source_async(source))
    # CorpusPipeline is synchronous. If an async host invokes it, isolate the
    # MinerU coroutine in a worker thread instead of nesting asyncio.run().
    with ThreadPoolExecutor(max_workers=1, thread_name_prefix="mineru-parser") as executor:
        return executor.submit(lambda: asyncio.run(parse_source_async(source))).result()
